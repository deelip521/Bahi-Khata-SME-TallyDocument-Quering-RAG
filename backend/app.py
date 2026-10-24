import re
import uuid
from contextlib import asynccontextmanager
from pathlib import Path

from fastapi import FastAPI, Depends, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from pypdf import PdfReader
from sqlalchemy import create_engine, String, Text, select
from sqlalchemy.orm import (
    DeclarativeBase,
    Mapped,
    mapped_column,
    Session,
    sessionmaker,
)

# --------------------------------------------------
# DATABASE AND STORAGE
# --------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent
DATABASE_URL = f"sqlite:///{BASE_DIR / 'bahi_khata.db'}"
UPLOAD_DIR = BASE_DIR / "uploads"

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
ALLOWED_EXTENSIONS = {".pdf", ".xlsx", ".txt"}

engine = create_engine(
    DATABASE_URL,
    connect_args={"check_same_thread": False},
)
SessionLocal = sessionmaker(
    bind=engine,
    expire_on_commit=False,
)


class Base(DeclarativeBase):
    pass


class Company(Base):
    __tablename__ = "companies"

    id: Mapped[str] = mapped_column(
        String(36), primary_key=True
    )
    name: Mapped[str] = mapped_column(
        String(150), nullable=False
    )
    email: Mapped[str] = mapped_column(
        String(254), unique=True, index=True, nullable=False
    )
    industry: Mapped[str] = mapped_column(
        String(100), nullable=False
    )


class Document(Base):
    __tablename__ = "documents"

    id: Mapped[str] = mapped_column(
        String(36), primary_key=True
    )
    company_id: Mapped[str] = mapped_column(
        String(36), nullable=False, index=True
    )
    filename: Mapped[str] = mapped_column(
        String(255), nullable=False
    )
    content_type: Mapped[str] = mapped_column(
        String(100), nullable=False
    )
    file_path: Mapped[str] = mapped_column(
        String(500), nullable=False
    )
    extracted_text: Mapped[str] = mapped_column(
        Text, nullable=False, default=""
    )


class CompanyCreate(BaseModel):
    name: str = Field(min_length=2, max_length=150)
    email: str = Field(min_length=5, max_length=254)
    industry: str = Field(min_length=2, max_length=100)


class AskRequest(BaseModel):
    company_id: str
    question: str = Field(min_length=2, max_length=1000)


def get_db():
    with SessionLocal() as db:
        yield db


@asynccontextmanager
async def lifespan(app: FastAPI):
    Base.metadata.create_all(bind=engine)
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    yield


app = FastAPI(
    title="Bahi-Khata API",
    description="Multi-company accounting document search API",
    version="1.1.0",
    lifespan=lifespan,
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:3000",
        "http://127.0.0.1:3000",
        "http://localhost:5173",
        "http://127.0.0.1:5173",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# --------------------------------------------------
# TEXT EXTRACTION
# --------------------------------------------------

def extract_text(file_path: Path, extension: str) -> str:
    if extension == ".pdf":
        reader = PdfReader(str(file_path))
        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    if extension == ".xlsx":
        workbook = load_workbook(
            file_path, read_only=True, data_only=True
        )
        lines = []

        try:
            for sheet in workbook.worksheets:
                lines.append(f"Sheet: {sheet.title}")

                for row in sheet.iter_rows(values_only=True):
                    lines.append(
                        " | ".join(
                            str(value) if value is not None else ""
                            for value in row
                        )
                    )
        finally:
            workbook.close()

        return "\n".join(lines)

    return file_path.read_text(
        encoding="utf-8",
        errors="replace",
    )


# --------------------------------------------------
# BASIC ROUTES
# --------------------------------------------------

@app.get("/")
def home():
    return {
        "message": "Bahi-Khata Python backend is running!",
        "status": "success",
    }


@app.get("/health")
def health_check():
    return {
        "status": "healthy",
        "service": "bahi-khata-backend",
    }


# --------------------------------------------------
# COMPANY REGISTRATION AND LISTING
# --------------------------------------------------

@app.post("/companies", status_code=201)
def register_company(
    company_data: CompanyCreate,
    db: Session = Depends(get_db),
):
    name = company_data.name.strip()
    email = company_data.email.strip().lower()
    industry = company_data.industry.strip()

    if len(name) < 2 or len(industry) < 2:
        raise HTTPException(
            status_code=422,
            detail="Company name and industry are required.",
        )

    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        raise HTTPException(
            status_code=422,
            detail="Enter a valid email address.",
        )

    existing = db.scalar(
        select(Company).where(Company.email == email)
    )

    if existing:
        raise HTTPException(
            status_code=409,
            detail="A company with this email is already registered.",
        )

    company = Company(
        id=str(uuid.uuid4()),
        name=name,
        email=email,
        industry=industry,
    )

    db.add(company)
    db.commit()
    db.refresh(company)

    return {
        "message": "Company registered successfully",
        "company": {
            "id": company.id,
            "name": company.name,
            "email": company.email,
            "industry": company.industry,
        },
    }


@app.get("/companies")
def list_companies(db: Session = Depends(get_db)):
    companies = db.scalars(
        select(Company).order_by(Company.name)
    ).all()

    return {
        "companies": [
            {
                "id": company.id,
                "name": company.name,
                "email": company.email,
                "industry": company.industry,
            }
            for company in companies
        ]
    }


# --------------------------------------------------
# UPLOAD DOCUMENTS
# --------------------------------------------------

@app.post("/companies/{company_id}/documents", status_code=201)
async def upload_document(
    company_id: str,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    company = db.get(Company, company_id)

    if company is None:
        raise HTTPException(
            status_code=404,
            detail="Company not found.",
        )

    filename = Path(file.filename or "upload").name
    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=415,
            detail="Only PDF, XLSX, and TXT files are allowed.",
        )

    content = await file.read(MAX_FILE_SIZE + 1)

    if not content:
        await file.close()
        raise HTTPException(
            status_code=400,
            detail="The file is empty.",
        )

    if len(content) > MAX_FILE_SIZE:
        await file.close()
        raise HTTPException(
            status_code=413,
            detail="Maximum file size is 10 MB.",
        )

    # Each company's files are stored in its own directory.
    storage_dir = UPLOAD_DIR / company_id
    storage_dir.mkdir(parents=True, exist_ok=True)

    document_id = str(uuid.uuid4())
    stored_path = storage_dir / f"{document_id}{extension}"

    try:
        stored_path.write_bytes(content)
        extracted_text = extract_text(stored_path, extension)

    except Exception as exc:
        stored_path.unlink(missing_ok=True)
        raise HTTPException(
            status_code=422,
            detail=(
                "Could not read this file. Please check that it is "
                "a valid PDF, XLSX, or TXT file."
            ),
        ) from exc

    finally:
        await file.close()

    document = Document(
        id=document_id,
        company_id=company_id,
        filename=filename,
        content_type=file.content_type or "application/octet-stream",
        file_path=str(stored_path),
        extracted_text=extracted_text,
    )

    try:
        db.add(document)
        db.commit()
    except Exception:
        db.rollback()
        stored_path.unlink(missing_ok=True)
        raise HTTPException(
            status_code=500,
            detail="Could not save document metadata.",
        )

    return {
        "message": "Document uploaded successfully",
        "document_id": document_id,
        "company_id": company_id,
        "filename": filename,
        "extracted_characters": len(extracted_text),
        "text_preview": extracted_text[:500],
    }


# --------------------------------------------------
# LIST A COMPANY'S DOCUMENTS
# --------------------------------------------------

@app.get("/companies/{company_id}/documents")
def list_documents(
    company_id: str,
    db: Session = Depends(get_db),
):
    company = db.get(Company, company_id)

    if company is None:
        raise HTTPException(
            status_code=404,
            detail="Company not found.",
        )

    documents = db.scalars(
        select(Document)
        .where(Document.company_id == company_id)
        .order_by(Document.filename)
    ).all()

    return {
        "company_id": company_id,
        "documents": [
            {
                "id": document.id,
                "filename": document.filename,
                "content_type": document.content_type,
            }
            for document in documents
        ],
    }


# --------------------------------------------------
# ASK QUESTIONS ABOUT A COMPANY'S DOCUMENTS
# Basic keyword search; not an AI model yet.
# --------------------------------------------------

@app.post("/ask")
def ask_question(
    request: AskRequest,
    db: Session = Depends(get_db),
):
    company = db.get(Company, request.company_id)

    if company is None:
        raise HTTPException(
            status_code=404,
            detail="Company not found.",
        )

    documents = db.scalars(
        select(Document).where(
            Document.company_id == request.company_id
        )
    ).all()

    if not documents:
        raise HTTPException(
            status_code=404,
            detail="No documents uploaded for this company yet.",
        )

    question = request.question.strip()
    question_words = {
        word.lower()
        for word in re.findall(r"[a-zA-Z0-9.]+", question)
        if len(word) > 1
    }

    matches = []

    for document in documents:
        text = document.extracted_text or ""
        lines = text.splitlines()

        for line_number, line in enumerate(lines, start=1):
            line = line.strip()

            if not line:
                continue

            lower_line = line.lower()
            score = sum(
                1 for word in question_words
                if word in lower_line
            )

            # Amount-related questions should also find lines
            # containing totals, amounts, invoices, or payments.
            asks_about_amount = any(
                term in question.lower()
                for term in ("amount", "total", "cost", "price", "paid")
            )

            if asks_about_amount and any(
                term in lower_line
                for term in (
                    "total", "amount", "invoice", "payment",
                    "paid", "balance", "₹", "rs.", "inr"
                )
            ):
                score += 2

            if score > 0:
                matches.append({
                    "filename": document.filename,
                    "line": line_number,
                    "text": line[:1000],
                    "score": score,
                })

    matches.sort(key=lambda item: item["score"], reverse=True)
    matches = matches[:10]

    if not matches:
        return {
            "company_id": request.company_id,
            "question": question,
            "answer": (
                "I couldn't find a matching passage in this company's "
                "uploaded documents. Try a more specific question."
            ),
            "sources": [],
            "search_type": "keyword",
        }

    answer = "\n".join(
        f"{item['filename']}, line {item['line']}: {item['text']}"
        for item in matches
    )

    return {
        "company_id": request.company_id,
        "question": question,
        "answer": answer,
        "sources": [
            {
                "filename": item["filename"],
                "line": item["line"],
                "text": item["text"],
            }
            for item in matches
        ],
        "search_type": "keyword",
        "note": (
            "This is keyword-based retrieval, not AI-generated reasoning. "
            "Verify financial figures against the original document."
        ),
    }